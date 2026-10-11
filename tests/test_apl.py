import importlib.util
import json
from pathlib import Path
import re
import tempfile
import unittest
import zipfile

import docx
from docx.oxml.ns import qn

SKILL = Path(__file__).resolve().parents[1] / 'plugins/indian-law/skills/quashing-application'
SPEC = importlib.util.spec_from_file_location('build_apl', SKILL / 'scripts/build_apl.py')
build_apl = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(build_apl)
EXAMPLE = json.loads((SKILL / 'assets/example-apl.json').read_text())


class AplBuilderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.out = Path(self.tmp.name) / 'matter' / 'APL.docx'

    def render(self, spec=EXAMPLE):
        report = build_apl.build(spec, self.out)
        return report, docx.Document(str(self.out))

    def test_house_format_is_preserved(self):
        _, document = self.render()
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.inches, 8.27, places=1)
        self.assertAlmostEqual(section.top_margin.inches, 1.97, places=1)
        self.assertAlmostEqual(section.left_margin.inches, 1.97, places=1)
        self.assertAlmostEqual(section.bottom_margin.inches, 1.0, places=1)
        self.assertAlmostEqual(section.right_margin.inches, 1.0, places=1)
        paragraphs = document.paragraphs + [p for c in document.tables[0]._cells for p in c.paragraphs]
        runs = [r for p in paragraphs for r in p.runs]
        self.assertTrue(runs)
        self.assertEqual({r.font.name for r in runs}, {'Times New Roman'})
        self.assertEqual({r.font.size.pt for r in runs}, {14.0})
        self.assertEqual({p.paragraph_format.line_spacing for p in document.paragraphs}, {2.0})
        borders = document.tables[0]._tbl.tblPr.find(qn('w:tblBorders'))
        self.assertIsNotNone(borders, 'cause title must stay boxed')
        self.assertIn('PAGE', ''.join(section.header._element.itertext()) + section.header._element.xml)

    def test_numbering_counts_and_report(self):
        report, document = self.render()
        self.assertEqual(report['numbered_paragraphs'], 7)
        self.assertEqual(report['grounds'], 1)
        self.assertEqual(report['annexures_referenced'], ['ANNEXURE A'])
        self.assertTrue(report['unresolved_placeholders'])
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('Paras 1 to 7', text)
        self.assertNotIn('[[', text)
        self.assertIn('GROUNDS', text)
        self.assertIn('Hence, this Application.', text)

    def test_markup_bold_and_underline(self):
        _, document = self.render()
        heading = document.paragraphs[0]
        self.assertTrue(all(r.bold and r.underline for r in heading.runs))
        cell = document.tables[0].rows[0].cells[1].paragraphs[0]
        self.assertTrue(cell.runs[0].bold)

    def test_refuses_to_write_inside_plugin(self):
        with self.assertRaises(build_apl.SpecError):
            build_apl.build(EXAMPLE, SKILL / 'assets' / 'leak.docx')
        self.assertFalse((SKILL / 'assets' / 'leak.docx').exists())

    def test_rejects_incomplete_spec(self):
        for key in ('applicants', 'facts', 'grounds', 'prayers'):
            spec = dict(EXAMPLE)
            spec.pop(key)
            with self.subTest(missing=key), self.assertRaises(build_apl.SpecError):
                build_apl.build(spec, self.out)

    def test_template_contains_only_placeholder_text(self):
        allowed = re.compile(r'^(\[\[[a-z_]+\]\].*|APPLICANTS|NON-APPLICANTS|::\s*VERSUS ::|\s*|\d+)$')
        with zipfile.ZipFile(SKILL / 'assets/apl-template.docx') as package:
            for name in package.namelist():
                if not name.endswith('.xml'):
                    continue
                xml = package.read(name).decode('utf-8')
                for para in re.findall(r'<w:p[ >].*?</w:p>', xml, flags=re.S):
                    text = ''.join(re.findall(r'<w:t[^>]*>([^<]*)', para)).strip()
                    with self.subTest(part=name, text=text[:40]):
                        self.assertRegex(text, allowed)
            core = package.read('docProps/core.xml').decode('utf-8')
            self.assertRegex(core, r'<dc:creator\s*/>|<dc:creator></dc:creator>')


if __name__ == '__main__':
    unittest.main()
