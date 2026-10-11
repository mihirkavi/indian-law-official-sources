# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx==1.2.0"]
# ///
"""Render a Bombay High Court criminal application (APL) .docx from a JSON matter spec.

The bundled template carries the house formatting (A4, 2-inch top/left and 1-inch
bottom/right margins, Times New Roman 14 pt, double spacing, boxed cause title,
numbered paragraphs, lettered grounds and prayers, page number in the header).
Each formatting role is a prototype paragraph marked ``[[role]]`` in the template;
the builder clones those paragraph properties and writes the matter text.

Inline markup in any text: **bold**, ^^bold underlined^^, and a literal tab (\\t).
The output path is explicit; keep it inside the authorised matter workspace.
"""
import argparse
import copy
import json
from pathlib import Path
import re
import sys

import docx
from docx.oxml.ns import qn
from lxml import etree

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / 'assets' / 'apl-template.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MARK = re.compile(r'^\[\[([a-z_]+)\]\]')
TOKEN = re.compile(r'(\^\^.+?\^\^|\*\*.+?\*\*)')
PLACEHOLDER = re.compile(r'\[(?:FACT REQUIRED|TO BE [A-Z ]+)[^\]]*\]')
DEFAULT_HEADING = 'IN THE HON’BLE HIGH COURT OF JUDICATURE AT BOMBAY, NAGPUR BENCH, NAGPUR'
BODY_ROLES = {'heading', 'case_no', 'title', 'title_line', 'intro', 'num', 'num_flat', 'quote',
              'gap', 'grounds_h', 'ground', 'num_late', 'hence', 'prayer_h', 'prayer_first',
              'prayer', 'sign', 'blank', 'sa_h', 'sa', 'deponent', 'left'}
CELL_ROLES = {'party_num', 'party', 'na', 'na_ind'}


class SpecError(ValueError):
    pass


def _ppr(paragraph):
    pp = paragraph._p.pPr
    pp = copy.deepcopy(pp) if pp is not None else etree.Element(qn('w:pPr'))
    for r in pp.findall(qn('w:rPr')):
        pp.remove(r)
    return pp


def _rpr(bold=False, underline=False):
    return etree.fromstring(
        f'<w:rPr xmlns:w="{W}"><w:rFonts w:ascii="Times New Roman" w:eastAsia="Times New Roman" '
        'w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
        + ('<w:b/><w:bCs/>' if bold else '')
        + '<w:color w:val="000000"/><w:sz w:val="28"/><w:szCs w:val="28"/>'
        + ('<w:u w:val="single"/>' if underline else '') + '</w:rPr>')


def paragraph(ppr, text):
    p = etree.Element(qn('w:p'))
    pp = copy.deepcopy(ppr)
    pp.append(_rpr())
    p.append(pp)
    for part in TOKEN.split(text or ''):
        if not part:
            continue
        bold = underline = False
        if part.startswith('^^'):
            part, bold, underline = part[2:-2], True, True
        elif part.startswith('**'):
            part, bold = part[2:-2], True
        for k, seg in enumerate(part.split('\t')):
            if k:
                r = etree.SubElement(p, qn('w:r'))
                r.append(_rpr(bold, underline))
                etree.SubElement(r, qn('w:tab'))
            if seg:
                r = etree.SubElement(p, qn('w:r'))
                r.append(_rpr(bold, underline))
                t = etree.SubElement(r, qn('w:t'))
                t.text = seg
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return p


def load_prototypes(document):
    body, cells = {}, {}
    for p in document.paragraphs:
        m = MARK.match(p.text)
        if m:
            body[m.group(1)] = _ppr(p)
    table = document.tables[0]
    for c in table._cells:
        for p in c.paragraphs:
            m = MARK.match(p.text)
            if m:
                cells[m.group(1)] = _ppr(p)
    missing = (BODY_ROLES - body.keys()) | (CELL_ROLES - cells.keys())
    if missing:
        raise SpecError(f'Template is missing prototype roles: {sorted(missing)}')
    return body, cells, table


def _require(spec, key, kind):
    value = spec.get(key)
    if not isinstance(value, kind) or not value:
        raise SpecError(f'Spec field "{key}" is required ({kind.__name__})')
    return value


def _parties(entries, label):
    if not isinstance(entries, list) or not entries:
        raise SpecError(f'At least one {label} is required')
    for e in entries:
        if not isinstance(e, dict) or not isinstance(e.get('name'), str) or not e['name'].strip():
            raise SpecError(f'Each {label} needs a "name"')
    return entries


def build(spec, output, template=TEMPLATE):
    output = Path(output).resolve()
    if HERE.parent in output.parents:
        raise SpecError('Write matter documents to the matter workspace, not inside the plugin')
    applicants = _parties(spec.get('applicants'), 'applicant')
    non_applicants = _parties(spec.get('non_applicants'), 'non-applicant')
    facts = _require(spec, 'facts', list)
    grounds = _require(spec, 'grounds', list)
    prayers = _require(spec, 'prayers', list)

    document = docx.Document(str(template))
    proto, cproto, table = load_prototypes(document)
    body = document.element.body
    sect = body.find(qn('w:sectPr'))
    tbl = copy.deepcopy(table._tbl)
    for el in list(body):
        if el is not sect:
            body.remove(el)

    def add(role, text):
        sect.addprevious(paragraph(proto[role], text))

    def refill(tc, items):
        for p in tc.findall(qn('w:p')):
            tc.remove(p)
        for role, text in items:
            tc.append(paragraph(cproto[role], text))

    app_items = []
    for a in applicants:
        app_items.append(('party_num', f"**{a['name']},** " + ' '.join(a.get('details', []))))
        app_items += [('party', line) for line in a.get('lines', [])]
        if a.get('relation'):
            app_items.append(('party', f"                         …..({a['relation']})"))
        app_items += [('party', line) for line in a.get('after', [])]
    na_items = []
    for i, n in enumerate(non_applicants, 1):
        na_items.append(('na', f"**{i}.  {n['name']},**"))
        na_items += [('na_ind', line) for line in n.get('lines', [])]
        if n.get('relation'):
            na_items.append(('na_ind', f"                         …..({n['relation']})"))
    rows = tbl.findall(qn('w:tr'))
    refill(rows[0].findall(qn('w:tc'))[1], app_items)
    refill(rows[-1].findall(qn('w:tc'))[1], na_items)

    add('heading', f"^^{spec.get('court_heading', DEFAULT_HEADING)}^^")
    add('case_no', f"^^{spec.get('case_number_line', 'CRIMINAL APPLICATION (APL) NO.             /20__')}^^")
    sect.addprevious(tbl)
    add('title', f"^^{spec.get('statutory_title', 'APPLICATION U/S 528 OF BHARATIYA NAGARIK SURAKSHA SANHITA, 2023')}^^")
    for line in spec.get('title_lines', []):
        add('title_line', line)
    add('intro', '            The Applicants above-named most humbly and respectfully \tbeg to submit as under: -')

    count = 0
    for item in facts:
        if isinstance(item, str):
            item = {'text': item}
        if 'quote' in item:
            add('quote', '          “' + item['quote'] + '”')
        else:
            count += 1
            add('num_flat' if item.get('flat') else 'num', item['text'])
    add('gap', '')
    add('grounds_h', '  ^^GROUNDS^^')
    for g in grounds:
        add('ground', g)
    for c in spec.get('concluding', []):
        count += 1
        add('num_late', c)
    add('hence', 'Hence, this Application.')
    add('prayer_h', "^^PRAYER^^**:-**  It is, therefore, most humbly prayed that this Hon'ble Court may kindly be pleased to:-")
    add('prayer_first', prayers[0])
    for p in prayers[1:]:
        add('prayer', p)
    add('sign', f"{spec.get('place', 'NAGPUR')}                                               {spec.get('counsel_label', 'Counsel For Applicants')}")
    add('sign', f"DATE- {spec.get('date', '    /    /20__')}  \t\t \t{spec.get('counsel', '[COUNSEL NAME]')}")
    add('blank', '')
    add('blank', '')
    add('sa_h', '^^SOLEMN AFFIRMATION^^')
    add('sa', '                 ' + spec.get('affirmation', '[FACT REQUIRED: solemn affirmation]').replace('{last_para}', str(count)))
    add('deponent', 'DEPONENT')
    add('left', 'I KNOW AND IDENTIFY THE DEPONENT')
    add('left', 'ADVOCATE \t\t')

    props = document.core_properties
    props.title = spec.get('document_title', 'Criminal Application (APL)')
    props.author = spec.get('counsel', '')
    props.last_modified_by = ''
    props.comments = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))

    text = '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(
        c.text for t in document.tables for c in t._cells)
    return {
        'output': str(output),
        'numbered_paragraphs': count,
        'grounds': len(grounds),
        'prayers': len(prayers),
        'unresolved_placeholders': PLACEHOLDER.findall(text),
        'annexures_referenced': sorted(set(re.findall(r'ANNEXURE [A-Z]+', text))),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('spec', type=Path, help='Matter spec JSON (see assets/example-apl.json)')
    parser.add_argument('--out', type=Path, required=True, help='Output .docx in the matter workspace')
    parser.add_argument('--template', type=Path, default=TEMPLATE)
    args = parser.parse_args()
    try:
        spec = json.loads(args.spec.read_text(encoding='utf-8'))
        if not isinstance(spec, dict):
            raise SpecError('Spec must be a JSON object')
        report = build(spec, args.out, args.template)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({'error': str(exc)}), file=sys.stderr)
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
